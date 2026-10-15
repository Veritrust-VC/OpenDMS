"""Lietu klasifikācijas shēma (nomenklatūra) — MK 282 31.–45. p.; kisc/dokumentu-arhivesanas-standarts 02.

Tīrās funkcijas (bez DB): parse_retention, detect_columns, rows_to_entries, parse_ark_xml, to_ark_xml.
DB funkcijas: settings, shēmu saraksts, koks, statusi, imports, aktīvā shēma.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Optional

from lxml import etree

from opendms.database import get_pool

ARK = "urn:lv:kisc:arhivesana:1.0"
XSD_PATH = Path(__file__).parent / "schemas" / "arhivesana-v1.xsd"

STATUSES = ["Projekts", "IzskatitsEK", "IesniegtsLNA", "SaskanotsLNA", "NoraiditsLNA",
            "NavJasaskano", "Apstiprinats", "Aizstata"]
# P1 (kisc/dokumentu-arhivesana): atļautās pārejas
TRANSITIONS = {
    "Projekts": {"IzskatitsEK"},
    "IzskatitsEK": {"IesniegtsLNA", "NavJasaskano", "Projekts"},
    "IesniegtsLNA": {"SaskanotsLNA", "NoraiditsLNA"},
    "NoraiditsLNA": {"Projekts"},
    "SaskanotsLNA": {"Apstiprinats"},
    "NavJasaskano": {"Apstiprinats"},
    "Apstiprinats": {"Aizstata"},
    "Aizstata": set(),
}
EDITABLE = {"Projekts", "NoraiditsLNA"}
STATUS_LV = {"Projekts": "Projekts", "IzskatitsEK": "Izskatīts EK", "IesniegtsLNA": "Iesniegts LNA",
             "SaskanotsLNA": "Saskaņots ar LNA", "NoraiditsLNA": "LNA noraidīts", "NavJasaskano": "Nav jāsaskaņo",
             "Apstiprinats": "Apstiprināts", "Aizstata": "Aizstāta"}


class SchemeError(Exception):
    pass


# ───────────────────────── glabāšanas termiņa teksts ─────────────────────────

_PERM = re.compile(r"^\s*(pastāvīg|pastavig|past\.|p\s*$)", re.I)
_YEARS = re.compile(r"(\d{1,3})\s*(g\b|g\.|gad|gads|gadi|gadus|y)", re.I)
_BARE = re.compile(r"^\s*(\d{1,3})\s*$")


def parse_retention(text: Any) -> dict:
    """'Pastāvīgi' | '10 g.' | '5 g. EK' | '75 g. AI' | 10 → {kind, years, marks, ok, warning}."""
    raw = "" if text is None else str(text).strip()
    marks = [m for m in ("EK", "AI") if re.search(rf"\b{m}\b", raw)]
    out = {"kind": None, "years": None, "marks": marks, "text": raw or None, "ok": False, "warning": None}
    if not raw:
        out["warning"] = "nav glabāšanas termiņa"
        return out
    if _PERM.search(raw):
        out.update(kind="Pastavigi", ok=True)
        return out
    m = _YEARS.search(raw) or _BARE.match(raw)
    if m:
        y = int(m.group(1))
        if y > 0:
            out.update(kind="Gadi", years=y, ok=True)
            return out
    out["warning"] = f"termiņu '{raw}' nevar nolasīt — norādiet gadus vai 'Pastāvīgi'"
    return out


# ───────────────────────── tabulas imports ─────────────────────────

COLUMNS = {
    "idx": ("indekss", "lietas indekss", "index", "nr. p. k", "kods"),
    "name": ("lietas nosaukums", "nosaukums", "lieta", "līmeņa nosaukums"),
    "retention": ("glabāšanas termiņš", "termiņš", "glabāšanas", "termins"),
    "legal_basis": ("pants", "pamatojums", "normatīv", "tiesiskais"),
    "responsible": ("atbildīg", "struktūrvienīb", "strukturvienib"),
    "medium": ("nesējs", "informācijas vide", "vide", "nesejs"),
    "system": ("informācijas sistēm", "sistēma", " is"),
    "restriction": ("ierobežojum", "pieejamīb"),
    "notes": ("piezīmes", "piezimes", "komentār"),
}


def detect_columns(header: list[Any]) -> dict[str, int]:
    """Kolonnu nosaukumi → lauki. Atgriež {lauks: kolonnas indekss}."""
    found: dict[str, int] = {}
    cells = [str(h or "").strip().lower() for h in header]
    for field, keys in COLUMNS.items():
        for i, c in enumerate(cells):
            if i in found.values() or not c:
                continue
            if any(k in c for k in keys):
                found[field] = i
                break
    return found


def _norm_idx(v: Any) -> str:
    s = str(v or "").strip()
    if re.fullmatch(r"\d+\.0", s):          # 1.0 no Excel skaitļa
        s = s[:-2]
    return s.rstrip(".")


def _idx_parts(idx: str) -> list[str]:
    return [p for p in re.split(r"[.\-]", idx) if p]


def _medium(v: Any) -> Optional[str]:
    s = str(v or "").strip().lower()
    if not s:
        return None
    if s in ("e", "elektronisks", "elektroniska", "el.") or s.startswith("elektron"):
        return "E"
    if s in ("p", "papīrs", "papīra") or s.startswith("papīr"):
        return "P"
    if s.startswith("hibr") or s == "h":
        return "H"
    return "C"


def rows_to_entries(rows: list[list[Any]]) -> dict:
    """Izklājlapas rindas → shēmas ieraksti. Rinda ar termiņu ir lieta; bez termiņa — līmenis.
    Līmeņa dziļums no indeksa (1 | 1.2 | 1.2.3); lietas vecāks — garākais līmeņa indekss, kas ir tās prefikss."""
    header_row = None
    for i, r in enumerate(rows[:20]):
        cols = detect_columns(r)
        if "idx" in cols and "name" in cols:
            header_row, mapping = i, cols
            break
    if header_row is None:
        raise SchemeError("Neatradu galveni ar ailēm 'Indekss' un 'Nosaukums' (pirmajās 20 rindās)")
    entries, warnings, seen = [], [], set()
    get = lambda r, f: (r[mapping[f]] if f in mapping and mapping[f] < len(r) else None)
    for n, r in enumerate(rows[header_row + 1:], start=header_row + 2):
        idx, name = _norm_idx(get(r, "idx")), str(get(r, "name") or "").strip()
        if not idx and not name:
            continue
        if not idx or not name:
            warnings.append(f"{n}. rinda: trūkst indeksa vai nosaukuma — izlaista")
            continue
        if idx in seen:
            warnings.append(f"{n}. rinda: indekss {idx} atkārtojas — izlaista")
            continue
        seen.add(idx)
        ret_raw = get(r, "retention")
        is_case = ret_raw not in (None, "") and str(ret_raw).strip() != ""
        e = {"row": n, "idx": idx, "name": name, "kind": "lieta" if is_case else "limenis",
             "level": min(len(_idx_parts(idx)), 3), "notes": (str(get(r, "notes")).strip() if get(r, "notes") else None),
             "responsible_unit": (str(get(r, "responsible")).strip() if get(r, "responsible") else None),
             "medium": _medium(get(r, "medium")), "information_system": (str(get(r, "system")).strip() if get(r, "system") else None),
             "retention_legal_basis": (str(get(r, "legal_basis")).strip() if get(r, "legal_basis") else None),
             "access_basis": (str(get(r, "restriction")).strip() if get(r, "restriction") else None)}
        if is_case:
            p = parse_retention(ret_raw)
            e.update(retention_kind=p["kind"], retention_years=p["years"], retention_marks=p["marks"], retention_text=p["text"])
            if p["warning"]:
                warnings.append(f"{n}. rinda ({idx}): {p['warning']}")
            e["disposition"] = "NodotArhivam" if p["kind"] == "Pastavigi" else ("Parskatit" if "EK" in p["marks"] else "Iznicinat")
            e["retention_trigger"] = "NoLietasSlegsanas"
            if e["access_basis"]:
                e["default_info_status"] = "IerobezotasPieejamibas"
        entries.append(e)
    # vecāki
    levels = {e["idx"]: e for e in entries if e["kind"] == "limenis"}
    for e in entries:
        parts = _idx_parts(e["idx"])
        parent = None
        for k in range(len(parts) - (1 if e["kind"] == "limenis" else 0), 0, -1):
            cand = [l for l in levels.values() if _idx_parts(l["idx"]) == parts[:k] and l is not e]
            if cand:
                parent = cand[0]["idx"]
                break
        e["parent_idx"] = parent
        if e["kind"] == "lieta" and parent is None:
            warnings.append(f"{e['row']}. rinda ({e['idx']}): lietai nav līmeņa — tiks ievietota saknē")
    n_cases = sum(1 for e in entries if e["kind"] == "lieta")
    return {"entries": entries, "warnings": warnings, "columns": {k: rows[header_row][v] for k, v in mapping.items()},
            "summary": {"limeni": len(entries) - n_cases, "lietas": n_cases,
                        "pastavigi": sum(1 for e in entries if e.get("retention_kind") == "Pastavigi"),
                        "bez_termina": sum(1 for e in entries if e["kind"] == "lieta" and not e.get("retention_kind"))}}


def read_table(filename: str, data: bytes) -> list[list[Any]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=";,\t")
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    raise SchemeError("Atbalstītie formāti: .xlsx, .csv, .xml (ark:KlasifikacijasShema)")


# ───────────────────────── ark XML ─────────────────────────

_schema = None


def _xsd() -> etree.XMLSchema:
    global _schema
    if _schema is None:
        _schema = etree.XMLSchema(etree.parse(str(XSD_PATH)))
    return _schema


def parse_ark_xml(data: bytes) -> dict:
    doc = etree.fromstring(data)
    if not _xsd().validate(etree.ElementTree(doc)):
        raise SchemeError("XML neatbilst arhivesana-v1.xsd: " + "; ".join(e.message for e in list(_xsd().error_log)[:3]))
    ns = {"a": ARK}
    t = lambda el, p: (el.findtext(p, namespaces=ns) or None)
    entries = []

    def walk(el, parent_idx):
        for child in el:
            tag = etree.QName(child).localname
            if tag == "Limenis":
                idx = t(child, "a:Indekss")
                entries.append({"idx": idx, "name": t(child, "a:Nosaukums"), "kind": "limenis",
                                "level": int(child.get("numurs")), "parent_idx": parent_idx})
                walk(child, idx)
            elif tag == "LietasIeraksts":
                g = child.find("a:GlabasanasTermins", ns)
                pi = child.find("a:PiekluvesIerobezojums", ns)
                entries.append({
                    "idx": t(child, "a:Indekss"), "name": t(child, "a:Nosaukums"), "kind": "lieta", "level": None,
                    "parent_idx": parent_idx, "retention_kind": g.get("veids"),
                    "retention_years": int(g.text) if (g.text or "").strip() else None,
                    "retention_marks": (g.get("atzimes") or "").split(), "retention_legal_basis": g.get("tiesiskaisPamatojums"),
                    "retention_trigger": t(child, "a:GlabasanasTriggeris"), "disposition": t(child, "a:Riciba"),
                    "responsible_unit": t(child, "a:Atbildigais/a:Strukturvieniba"),
                    "responsible_person": t(child, "a:Atbildigais/a:Persona/a:VardsUzvards"),
                    "medium": t(child, "a:Nesejs"), "information_system": t(child, "a:InformacijasSistema"),
                    "default_info_status": t(child, "a:NoklusetaisInformacijasStatuss"),
                    "access_basis": pi.findtext("a:TiesiskaisPamatojums", namespaces=ns) if pi is not None else None,
                    "access_until": pi.findtext("a:IerobezojumsLidz", namespaces=ns) if pi is not None else None,
                    "long_running": t(child, "a:Ilgstosa") == "true",
                    "continues_from_year": int(t(child, "a:TurpinasNoPerioda")) if t(child, "a:TurpinasNoPerioda") else None,
                    "notes": t(child, "a:Piezimes")})
    for lim in doc.findall("a:Limenis", ns):
        walk_root = lim
        entries.append({"idx": t(walk_root, "a:Indekss"), "name": t(walk_root, "a:Nosaukums"), "kind": "limenis",
                        "level": int(walk_root.get("numurs")), "parent_idx": None})
        walk(walk_root, t(walk_root, "a:Indekss"))
    meta = {"period_from": t(doc, "a:Periods/a:No"), "period_to": t(doc, "a:Periods/a:Lidz"), "version": t(doc, "a:Versija")}
    n_cases = sum(1 for e in entries if e["kind"] == "lieta")
    return {"entries": entries, "warnings": [], "meta": meta,
            "summary": {"limeni": len(entries) - n_cases, "lietas": n_cases,
                        "pastavigi": sum(1 for e in entries if e.get("retention_kind") == "Pastavigi"), "bez_termina": 0}}


def to_ark_xml(scheme: dict, entries: list[dict], org: dict, settings: dict) -> bytes:
    A = lambda tag: f"{{{ARK}}}{tag}"
    root = etree.Element(A("KlasifikacijasShema"), nsmap={"ark": ARK})
    root.set("id", f"shema-{scheme['id']}")
    ie = etree.SubElement(root, A("Iestade"))
    etree.SubElement(ie, A("OrgKods")).text = org.get("std_code") or "00-0000"
    etree.SubElement(ie, A("Nosaukums")).text = org.get("name")
    if settings.get("registration_number"):
        etree.SubElement(ie, A("RegistracijasNr")).text = settings["registration_number"]
    p = etree.SubElement(root, A("Periods"))
    etree.SubElement(p, A("No")).text = str(scheme["period_from"])
    etree.SubElement(p, A("Lidz")).text = str(scheme["period_to"])
    etree.SubElement(root, A("Versija")).text = scheme.get("version") or "1"
    hist = scheme.get("history") or []
    for h in hist or [{"statuss": scheme["status"]}]:
        s = etree.SubElement(root, A("Saskanosana"))
        etree.SubElement(s, A("Statuss")).text = h["statuss"] if h["statuss"] != "Aizstata" else "Apstiprinats"
        if h.get("datums"):
            etree.SubElement(s, A("Datums")).text = h["datums"]
        if h.get("atsauce"):
            etree.SubElement(s, A("Atsauce")).text = h["atsauce"]
        if h.get("piezimes"):
            etree.SubElement(s, A("Piezimes")).text = h["piezimes"]
    by_parent: dict[Optional[int], list[dict]] = {}
    for e in entries:
        by_parent.setdefault(e.get("parent_id"), []).append(e)

    def emit_case(parent_el, e):
        c = etree.SubElement(parent_el, A("LietasIeraksts"))
        etree.SubElement(c, A("Indekss")).text = e["idx"]
        etree.SubElement(c, A("Nosaukums")).text = e["name"]
        g = etree.SubElement(c, A("GlabasanasTermins"))
        g.set("veids", e["retention_kind"])
        g.text = str(e["retention_years"]) if e["retention_kind"] == "Gadi" and e.get("retention_years") else ""
        if e.get("retention_marks"):
            g.set("atzimes", " ".join(e["retention_marks"]))
        if e.get("retention_legal_basis"):
            g.set("tiesiskaisPamatojums", e["retention_legal_basis"])
        etree.SubElement(c, A("GlabasanasTriggeris")).text = e.get("retention_trigger") or "NoLietasSlegsanas"
        etree.SubElement(c, A("Riciba")).text = e.get("disposition") or "Iznicinat"
        at = etree.SubElement(c, A("Atbildigais"))
        etree.SubElement(at, A("Strukturvieniba")).text = e.get("responsible_unit") or "Nav norādīta"
        if e.get("responsible_person"):
            ps = etree.SubElement(at, A("Persona"))
            etree.SubElement(ps, A("VardsUzvards")).text = e["responsible_person"]
        etree.SubElement(c, A("Nesejs")).text = e.get("medium") or "E"
        if e.get("information_system"):
            etree.SubElement(c, A("InformacijasSistema")).text = e["information_system"]
        if e.get("default_info_status"):
            etree.SubElement(c, A("NoklusetaisInformacijasStatuss")).text = e["default_info_status"]
        if e.get("access_basis"):
            pi = etree.SubElement(c, A("PiekluvesIerobezojums"))
            etree.SubElement(pi, A("TiesiskaisPamatojums")).text = e["access_basis"]
            if e.get("access_until"):
                etree.SubElement(pi, A("IerobezojumsLidz")).text = str(e["access_until"])
        if e.get("long_running"):
            etree.SubElement(c, A("Ilgstosa")).text = "true"
        if e.get("continues_from_year"):
            etree.SubElement(c, A("TurpinasNoPerioda")).text = str(e["continues_from_year"])
        if e.get("notes"):
            etree.SubElement(c, A("Piezimes")).text = e["notes"]

    def emit_level(parent_el, e, depth):
        l = etree.SubElement(parent_el, A("Limenis"))
        l.set("numurs", str(min(depth, 3)))
        etree.SubElement(l, A("Indekss")).text = e["idx"]
        etree.SubElement(l, A("Nosaukums")).text = e["name"]
        kids = by_parent.get(e["id"], [])
        if not kids:
            raise SchemeError(f"Līmenim {e['idx']} nav ne apakšlīmeņu, ne lietu")
        for k in kids:
            emit_level(l, k, depth + 1) if k["kind"] == "limenis" else emit_case(l, k)

    top = by_parent.get(None, [])
    loose = [e for e in top if e["kind"] == "lieta"]
    if loose:   # lietas bez līmeņa — XSD prasa līmeni
        l = etree.SubElement(root, A("Limenis"))
        l.set("numurs", "1")
        etree.SubElement(l, A("Indekss")).text = "0"
        etree.SubElement(l, A("Nosaukums")).text = "Lietas bez līmeņa"
        for e in loose:
            emit_case(l, e)
    for e in top:
        if e["kind"] == "limenis":
            emit_level(root, e, 1)
    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", pretty_print=True)
    if not _xsd().validate(etree.fromstring(data).getroottree()):
        raise SchemeError("Eksports neatbilst arhivesana-v1.xsd: " + "; ".join(e.message for e in list(_xsd().error_log)[:3]))
    return data


# ───────────────────────── DB ─────────────────────────

SETTINGS_FIELDS = ("archive_code", "archive_name", "fonds_number", "fonds_name", "registration_number",
                   "enforce_registration_minimum")
ENTRY_FIELDS = ("kind", "level", "idx", "name", "retention_kind", "retention_years", "retention_marks", "retention_text",
                "retention_legal_basis", "retention_trigger", "disposition", "responsible_unit", "responsible_person",
                "medium", "information_system", "default_info_status", "access_basis", "access_until", "long_running",
                "continues_from_year", "notes")


def _j(v):
    return json.loads(v) if isinstance(v, str) else v


async def get_settings(org: dict) -> dict:
    pool = await get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT * FROM archive_settings WHERE org_id = $1", org["id"])
    s = dict(row) if row else {"org_id": org["id"], "enforce_registration_minimum": False}
    if not s.get("registration_number") and org.get("reg_number"):
        s["registration_number_from_org"] = org["reg_number"]
    s["missing"] = [f for f in ("archive_code", "archive_name", "fonds_number", "fonds_name") if not s.get(f)]
    if not (s.get("registration_number") or org.get("reg_number")):
        s["missing"].append("registration_number")
    return s


async def put_settings(org: dict, data: dict, who: str) -> dict:
    vals = {k: data.get(k) for k in SETTINGS_FIELDS if k in data}
    if vals.get("enforce_registration_minimum"):
        if not await active_scheme(org["id"], date.today()):
            raise SchemeError("Reģistrācijas minimumu var ieslēgt tikai, kad šim periodam ir apstiprināta klasifikācijas shēma")
    pool = await get_pool()
    async with pool.acquire() as conn:
        await conn.execute("INSERT INTO archive_settings (org_id) VALUES ($1) ON CONFLICT DO NOTHING", org["id"])
        if vals:
            sets = ", ".join(f"{k} = ${i + 2}" for i, k in enumerate(vals))
            await conn.execute(f"UPDATE archive_settings SET {sets}, updated_at = NOW(), updated_by = ${len(vals) + 2} WHERE org_id = $1",
                               org["id"], *vals.values(), who)
    return await get_settings(org)


async def list_schemes(org_id: int) -> list[dict]:
    pool = await get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            """SELECT s.*, (SELECT count(*) FROM scheme_entries e WHERE e.scheme_id = s.id AND e.kind = 'lieta') AS lietas,
                      (SELECT count(*) FROM scheme_entries e WHERE e.scheme_id = s.id AND e.kind = 'lieta' AND e.retention_kind = 'Pastavigi') AS pastavigi
                 FROM classification_schemes s WHERE s.org_id = $1 ORDER BY period_from DESC, id DESC""", org_id)
    out = []
    for r in rows:
        d = dict(r)
        d["history"] = _j(d["history"])
        d["source"] = _j(d["source"])
        d["status_lv"] = STATUS_LV.get(d["status"])
        d["next"] = sorted(TRANSITIONS.get(d["status"], ()))
        d["editable"] = d["status"] in EDITABLE
        out.append(d)
    return out


async def get_scheme(org_id: int, scheme_id: int) -> tuple[dict, list[dict]]:
    pool = await get_pool()
    async with pool.acquire() as conn:
        s = await conn.fetchrow("SELECT * FROM classification_schemes WHERE id = $1 AND org_id = $2", scheme_id, org_id)
        if not s:
            raise SchemeError("Shēma nav atrasta")
        entries = [dict(r) for r in await conn.fetch(
            "SELECT * FROM scheme_entries WHERE scheme_id = $1 ORDER BY sort_order, id", scheme_id)]
    s = dict(s)
    s["history"], s["source"] = _j(s["history"]), _j(s["source"])
    s["status_lv"], s["next"], s["editable"] = STATUS_LV.get(s["status"]), sorted(TRANSITIONS.get(s["status"], ())), s["status"] in EDITABLE
    return s, entries


async def create_scheme(org_id: int, period_from: date, period_to: date, title: Optional[str], sector: bool,
                        copy_from: Optional[int], who: str) -> dict:
    if period_to < period_from:
        raise SchemeError("Perioda beigas ir pirms sākuma")
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            sid = await conn.fetchval(
                """INSERT INTO classification_schemes (org_id, period_from, period_to, title, sector, history)
                   VALUES ($1,$2,$3,$4,$5,$6) RETURNING id""",
                org_id, period_from, period_to, title, sector,
                json.dumps([{"statuss": "Projekts", "datums": date.today().isoformat(), "kas": who,
                             "piezimes": f"Izveidots no shēmas {copy_from}" if copy_from else "Izveidots"}], ensure_ascii=False))
            if copy_from:
                src = await conn.fetchrow("SELECT id, period_from FROM classification_schemes WHERE id = $1 AND org_id = $2", copy_from, org_id)
                if not src:
                    raise SchemeError("Avota shēma nav atrasta")
                rows = await conn.fetch("SELECT * FROM scheme_entries WHERE scheme_id = $1 ORDER BY sort_order, id", copy_from)
                idmap: dict[int, int] = {}
                for r in rows:
                    d = dict(r)
                    if d["kind"] == "lieta" and d["long_running"]:
                        d["continues_from_year"] = src["period_from"].year     # 43. p.: turpināmā lieta
                    new = await conn.fetchval(
                        f"""INSERT INTO scheme_entries (scheme_id, parent_id, sort_order, {", ".join(ENTRY_FIELDS)})
                            VALUES ($1, $2, $3, {", ".join(f"${i + 4}" for i in range(len(ENTRY_FIELDS)))}) RETURNING id""",
                        sid, idmap.get(d["parent_id"]), d["sort_order"], *[d[f] for f in ENTRY_FIELDS])
                    idmap[d["id"]] = new
                await conn.execute("UPDATE classification_schemes SET source = $2 WHERE id = $1", sid,
                                   json.dumps({"veids": "kopija", "no_shemas": copy_from}))
    return (await get_scheme(org_id, sid))[0]


async def replace_entries(org_id: int, scheme_id: int, parsed: dict, source: dict) -> dict:
    s, _ = await get_scheme(org_id, scheme_id)
    if not s["editable"]:
        raise SchemeError(f"Shēmu statusā '{STATUS_LV[s['status']]}' mainīt nevar")
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            await conn.execute("DELETE FROM scheme_entries WHERE scheme_id = $1", scheme_id)
            ids: dict[str, int] = {}
            # vecāki vispirms (līmeņi pēc dziļuma), secība avotā saglabājas sort_order
            ordered = sorted(enumerate(parsed["entries"]),
                             key=lambda p: (p[1]["kind"] == "lieta", len(_idx_parts(p[1]["idx"])), p[0]))
            for n, e in ordered:
                vals = {f: e.get(f) for f in ENTRY_FIELDS}
                vals["retention_marks"] = vals["retention_marks"] or []
                vals["long_running"] = bool(vals["long_running"])
                if isinstance(vals.get("access_until"), str):
                    vals["access_until"] = date.fromisoformat(vals["access_until"])
                if vals["kind"] == "lieta" and not vals["retention_kind"]:
                    # termiņš nav nolasāms — ieraksts paliek, bet lietā; apstiprināt neļaus (sk. set_status)
                    vals["retention_kind"], vals["retention_years"] = "Gadi", None
                    vals["notes"] = ((vals["notes"] or "") + " [TERMIŅŠ JĀPRECIZĒ]").strip()
                ids[e["idx"]] = await conn.fetchval(
                    f"""INSERT INTO scheme_entries (scheme_id, parent_id, sort_order, {", ".join(ENTRY_FIELDS)})
                        VALUES ($1, $2, $3, {", ".join(f"${i + 4}" for i in range(len(ENTRY_FIELDS)))}) RETURNING id""",
                    scheme_id, ids.get(e.get("parent_idx")), n, *vals.values())
            await conn.execute("UPDATE classification_schemes SET source = $2, updated_at = NOW() WHERE id = $1",
                               scheme_id, json.dumps(source, ensure_ascii=False))
    return await problems(org_id, scheme_id)


async def problems(org_id: int, scheme_id: int) -> dict:
    """Kas neļauj shēmu apstiprināt."""
    _, entries = await get_scheme(org_id, scheme_id)
    cases = [e for e in entries if e["kind"] == "lieta"]
    issues = []
    if not cases:
        issues.append("Shēmā nav nevienas lietas")
    for e in cases:
        if e["retention_kind"] == "Gadi" and not e["retention_years"]:
            issues.append(f"{e['idx']}: nav glabāšanas termiņa gados")
        if e["default_info_status"] in ("IerobezotasPieejamibas", "DienestaVajadzibam") and not e["access_basis"]:
            issues.append(f"{e['idx']}: ierobežotai informācijai nav tiesiskā pamatojuma")
    return {"lietas": len(cases), "limeni": len(entries) - len(cases), "problemas": issues}


async def set_status(org_id: int, scheme_id: int, new: str, when: Optional[str], ref: Optional[str], note: Optional[str], who: str) -> dict:
    s, _ = await get_scheme(org_id, scheme_id)
    if new not in TRANSITIONS.get(s["status"], set()):
        raise SchemeError(f"No statusa '{STATUS_LV[s['status']]}' uz '{STATUS_LV.get(new, new)}' pāriet nevar")
    if new in ("IzskatitsEK", "IesniegtsLNA", "Apstiprinats", "NavJasaskano"):
        p = await problems(org_id, scheme_id)
        if p["problemas"]:
            raise SchemeError("Shēmā ir nepilnības: " + "; ".join(p["problemas"][:5]))
    if new in ("SaskanotsLNA", "NoraiditsLNA") and not ref:
        raise SchemeError("Norādiet LNA lēmuma atsauci (numuru vai datumu)")
    pool = await get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            if new == "Apstiprinats":
                # 41. p.: viena shēma periodam — iepriekšējā, kas pārklājas, tiek aizstāta
                await conn.execute(
                    """UPDATE classification_schemes SET status = 'Aizstata', updated_at = NOW(),
                              history = history || $3::jsonb
                        WHERE org_id = $1 AND id <> $2 AND status = 'Apstiprinats'
                          AND daterange(period_from, period_to, '[]') && (SELECT daterange(period_from, period_to, '[]') FROM classification_schemes WHERE id = $2)""",
                    org_id, scheme_id, json.dumps([{"statuss": "Aizstata", "datums": date.today().isoformat(), "kas": who,
                                                    "piezimes": f"Aizstāta ar shēmu {scheme_id}"}], ensure_ascii=False))
            await conn.execute(
                """UPDATE classification_schemes SET status = $2, updated_at = NOW(), history = history || $3::jsonb WHERE id = $1""",
                scheme_id, new, json.dumps([{"statuss": new, "datums": when or date.today().isoformat(), "atsauce": ref,
                                             "piezimes": note, "kas": who,
                                             "kad": datetime.now(timezone.utc).isoformat()}], ensure_ascii=False))
    return (await get_scheme(org_id, scheme_id))[0]


async def delete_scheme(org_id: int, scheme_id: int) -> None:
    s, _ = await get_scheme(org_id, scheme_id)
    if s["status"] != "Projekts":
        raise SchemeError("Dzēst var tikai projektu")
    pool = await get_pool()
    async with pool.acquire() as conn:
        await conn.execute("DELETE FROM classification_schemes WHERE id = $1", scheme_id)


async def active_scheme(org_id: int, on: date) -> Optional[dict]:
    pool = await get_pool()
    async with pool.acquire() as conn:
        r = await conn.fetchrow(
            """SELECT * FROM classification_schemes WHERE org_id = $1 AND status = 'Apstiprinats'
                 AND $2 BETWEEN period_from AND period_to ORDER BY period_from DESC LIMIT 1""", org_id, on)
    return dict(r) if r else None
